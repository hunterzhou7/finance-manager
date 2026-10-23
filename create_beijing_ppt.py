#!/usr/bin/env python3
import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor

# Create presentation
prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

# Color definitions - spring/summer travel style
def set_shape_fill(shape, r, g, b):
    shape.fill.solid()
    shape.fill.fore_color.rgb = RGBColor(r, g, b)

def set_text_color(run, r, g, b):
    run.font.color.rgb = RGBColor(r, g, b)

# Image paths
img_dir = "/Volumes/lee/workspaces/claude_workspace/claude-my-product/010.内容运营（自媒体与运维）/beijing_images"
images = {
    "tiananmen": f"{img_dir}/beijing_tiananmen.png",
    "greatwall": f"{img_dir}/beijing_greatwall.png",
    "houhai": f"{img_dir}/beijing_houhai.png"
}

# Slide 1: Cover (天安门广场)
slide1 = prs.slides.add_slide(prs.slide_layouts[6])  # blank
# Background
bg1 = slide1.shapes.add_shape(1, Inches(0), Inches(0), prs.slide_width, prs.slide_height)
set_shape_fill(bg1, 25, 55, 100)  # Deep summer blue

# Title
title_box = slide1.shapes.add_textbox(Inches(0.5), Inches(2.5), Inches(12), Inches(1.5))
tf = title_box.text_frame
p = tf.paragraphs[0]
p.text = "北京3日游"
p.font.size = Pt(72)
p.font.bold = True
p.alignment = PP_ALIGN.CENTER
for run in p.runs:
    set_text_color(run, 255, 255, 255)

# Subtitle
sub_box = slide1.shapes.add_textbox(Inches(0.5), Inches(4), Inches(12), Inches(1))
tf = sub_box.text_frame
p = tf.paragraphs[0]
p.text = "春日漫游 · 探寻古今"
p.font.size = Pt(36)
p.alignment = PP_ALIGN.CENTER
for run in p.runs:
    set_text_color(run, 200, 220, 255)

# Image
if os.path.exists(images["tiananmen"]):
    slide1.shapes.add_picture(images["tiananmen"], Inches(4), Inches(5.2), width=Inches(5))

# Day marker
day1 = slide1.shapes.add_textbox(Inches(0.5), Inches(0.3), Inches(2), Inches(0.5))
tf = day1.text_frame
p = tf.paragraphs[0]
p.text = "DAY 1"
p.font.size = Pt(18)
for run in p.runs:
    set_text_color(run, 255, 183, 197)

# Slide 2: Great Wall (长城)
slide2 = prs.slides.add_slide(prs.slide_layouts[6])
bg2 = slide2.shapes.add_shape(1, Inches(0), Inches(0), prs.slide_width, prs.slide_height)
set_shape_fill(bg2, 34, 139, 103)  # Forest green summer

# Left content
content_box = slide2.shapes.add_textbox(Inches(0.5), Inches(2), Inches(5), Inches(4))
tf = content_box.text_frame
p = tf.paragraphs[0]
p.text = "八达岭长城"
p.font.size = Pt(44)
p.font.bold = True
for run in p.runs:
    set_text_color(run, 255, 255, 255)

p2 = tf.add_paragraph()
p2.text = "\n登长城做好了"
p2.font.size = Pt(28)
for run in p2.runs:
    set_text_color(run, 220, 255, 220)

p3 = tf.add_paragraph()
p3.text = "\n不到长城非好汉\n春天踏青必打卡"
p3.font.size = Pt(20)
for run in p3.runs:
    set_text_color(run, 180, 220, 180)

# Image
if os.path.exists(images["greatwall"]):
    slide2.shapes.add_picture(images["greatwall"], Inches(6), Inches(1.5), width=Inches(6.5))

# Day marker
day2 = slide2.shapes.add_textbox(Inches(0.5), Inches(0.3), Inches(2), Inches(0.5))
tf = day2.text_frame
p = tf.paragraphs[0]
p.text = "DAY 2"
p.font.size = Pt(18)
for run in p.runs:
    set_text_color(run, 170, 255, 200)

# Slide 3: Houhai (后海)
slide3 = prs.slides.add_slide(prs.slide_layouts[6])
bg3 = slide3.shapes.add_shape(1, Inches(0), Inches(0), prs.slide_width, prs.slide_height)
set_shape_fill(bg3, 255, 127, 80)  # Coral

# Left content
content_box3 = slide3.shapes.add_textbox(Inches(0.5), Inches(2), Inches(5), Inches(4))
tf = content_box3.text_frame
p = tf.paragraphs[0]
p.text = "后海酒吧街"
p.font.size = Pt(44)
p.font.bold = True
for run in p.runs:
    set_text_color(run, 255, 255, 255)

p4 = tf.add_paragraph()
p4.text = "\n夜幕降临 华灯初上"
p4.font.size = Pt(28)
for run in p4.runs:
    set_text_color(run, 255, 240, 200)

p5 = tf.add_paragraph()
p5.text = "\n胡同里的文艺气息\n春天的夜晚微凉却浪漫"
p5.font.size = Pt(20)
for run in p5.runs:
    set_text_color(run, 255, 220, 180)

# Image
if os.path.exists(images["houhai"]):
    slide3.shapes.add_picture(images["houhai"], Inches(6), Inches(1.5), width=Inches(6.5))

# Day marker
day3 = slide3.shapes.add_textbox(Inches(0.5), Inches(0.3), Inches(2), Inches(0.5))
tf = day3.text_frame
p = tf.paragraphs[0]
p.text = "DAY 3"
p.font.size = Pt(18)
for run in p.runs:
    set_text_color(run, 255, 200, 150)

# Save
output_dir = "/Volumes/lee/workspaces/claude_workspace/claude-my-product/010.内容运营（自媒体与运维）"
prs.save(f"{output_dir}/北京3日游.pptx")
print(f"PPT saved to {output_dir}/北京3日游.pptx")